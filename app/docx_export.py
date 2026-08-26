"""Markdown -> .docx without a pandoc dependency.

Covers what generated articles actually contain: headings, paragraphs, bullet
and numbered lists, blockquotes, fenced code, horizontal rules, and inline
bold/italic/code/links.
"""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

# inline: `code` | **bold** | __bold__ | *ital* | _ital_ | [text](url)
_INLINE = re.compile(
    r"(`[^`]+`)"
    r"|(\*\*[^*]+\*\*)"
    r"|(__[^_]+__)"
    r"|(\*[^*\n]+\*)"
    r"|(?<![A-Za-z0-9])(_[^_\n]+_)(?![A-Za-z0-9])"
    r"|(\[[^\]]+\]\([^)]+\))"
    # bare URL — generated reference lists are full of these, and they should
    # still be clickable once the piece lands in Word
    r"|(https?://[^\s<>\[\]()]+[^\s<>\[\]().,;:!?'\"])"
)
_LINK = re.compile(r"\[([^\]]+)\]\(([^)]+)\)")


def _add_hyperlink(paragraph, url: str, text: str) -> None:
    part = paragraph.part
    r_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "1155CC")
    rpr.append(color)
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    rpr.append(u)
    run.append(rpr)
    t = OxmlElement("w:t")
    t.text = text
    t.set(qn("xml:space"), "preserve")
    run.append(t)
    link.append(run)
    paragraph._p.append(link)


def _emit_inline(paragraph, text: str) -> None:
    pos = 0
    for m in _INLINE.finditer(text):
        if m.start() > pos:
            paragraph.add_run(text[pos:m.start()])
        tok = m.group(0)
        if tok.startswith("`"):
            r = paragraph.add_run(tok[1:-1])
            r.font.name = "Consolas"
            r.font.size = Pt(10)
            r.font.color.rgb = RGBColor(0xB0, 0x30, 0x60)
        elif tok.startswith("**") or tok.startswith("__"):
            paragraph.add_run(tok[2:-2]).bold = True
        elif tok.startswith("["):
            lm = _LINK.match(tok)
            if lm:
                _add_hyperlink(paragraph, lm.group(2), lm.group(1))
            else:
                paragraph.add_run(tok)
        elif tok.startswith(("http://", "https://")):
            _add_hyperlink(paragraph, tok, tok)
        else:
            paragraph.add_run(tok[1:-1]).italic = True
        pos = m.end()
    if pos < len(text):
        paragraph.add_run(text[pos:])


def _hr(doc: Document) -> None:
    p = doc.add_paragraph()
    pPr = p._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "BBBBBB")
    bdr.append(bottom)
    pPr.append(bdr)


def markdown_to_docx(md: str, out_path: Path, *, title: str = "",
                     subtitle: str = "") -> Path:
    doc = Document()

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(10)
    normal.paragraph_format.line_spacing = 1.15

    lines = md.replace("\r\n", "\n").split("\n")
    saw_h1 = False
    i = 0
    in_code = False
    code_buf: list[str] = []

    while i < len(lines):
        raw = lines[i]
        line = raw.rstrip()

        # fenced code
        if line.strip().startswith("```"):
            if in_code:
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Pt(18)
                p.paragraph_format.space_after = Pt(12)
                r = p.add_run("\n".join(code_buf))
                r.font.name = "Consolas"
                r.font.size = Pt(9.5)
                shd = OxmlElement("w:shd")
                shd.set(qn("w:val"), "clear")
                shd.set(qn("w:fill"), "F4F4F4")
                p._p.get_or_add_pPr().append(shd)
                code_buf, in_code = [], False
            else:
                in_code = True
            i += 1
            continue
        if in_code:
            code_buf.append(raw)
            i += 1
            continue

        if not line.strip():
            i += 1
            continue

        if re.fullmatch(r"\s*([-*_])\s*\1\s*\1[\s\-*_]*", line):
            _hr(doc)
            i += 1
            continue

        h = re.match(r"^(#{1,6})\s+(.*)$", line)
        if h:
            level = len(h.group(1))
            text = h.group(2).strip()
            if level == 1 and not saw_h1:
                saw_h1 = True
                p = doc.add_heading("", level=0)
                _emit_inline(p, text)
                if subtitle:
                    sp = doc.add_paragraph(subtitle)
                    sp.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    for r in sp.runs:
                        r.italic = True
                        r.font.size = Pt(10)
                        r.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
            else:
                p = doc.add_heading("", level=min(level, 4))
                _emit_inline(p, text)
            i += 1
            continue

        if line.lstrip().startswith(">"):
            body = re.sub(r"^\s*>\s?", "", line)
            p = doc.add_paragraph(style="Intense Quote")
            _emit_inline(p, body)
            i += 1
            continue

        b = re.match(r"^(\s*)[-*+]\s+(.*)$", line)
        if b:
            depth = min(len(b.group(1)) // 2, 2)
            style = "List Bullet" if depth == 0 else f"List Bullet {depth + 1}"
            p = doc.add_paragraph(style=_safe_style(doc, style, "List Bullet"))
            _emit_inline(p, b.group(2))
            i += 1
            continue

        n = re.match(r"^(\s*)\d+[.)]\s+(.*)$", line)
        if n:
            depth = min(len(n.group(1)) // 2, 2)
            style = "List Number" if depth == 0 else f"List Number {depth + 1}"
            p = doc.add_paragraph(style=_safe_style(doc, style, "List Number"))
            _emit_inline(p, n.group(2))
            i += 1
            continue

        # paragraph: join soft-wrapped lines
        buf = [line]
        j = i + 1
        while j < len(lines):
            nxt = lines[j].rstrip()
            if (not nxt.strip() or nxt.lstrip().startswith(("#", ">", "```"))
                    or re.match(r"^\s*([-*+]|\d+[.)])\s+", nxt)):
                break
            buf.append(nxt)
            j += 1
        p = doc.add_paragraph()
        _emit_inline(p, " ".join(s.strip() for s in buf))
        i = j

    if title and not saw_h1:
        h = doc.add_heading("", level=0)
        _emit_inline(h, title)
        # move the title to the top
        doc.paragraphs[0]._p.addprevious(h._p)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
    return out_path


def _safe_style(doc: Document, want: str, fallback: str) -> str:
    try:
        doc.styles[want]
        return want
    except KeyError:
        return fallback
