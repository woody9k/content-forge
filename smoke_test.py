"""End-to-end smoke test with the Claude CLI stubbed out.

    CF_FAKE_CLAUDE=1 CF_CONTENT_ROOT=/tmp/cf-test .venv/bin/python smoke_test.py

Exercises every route, then opens the generated .docx to confirm it parses and
that inline formatting survived the conversion.
"""
import json
import os
import shutil
import sys
import time
from pathlib import Path

os.environ.setdefault("CF_FAKE_CLAUDE", "1")
ROOT = Path(os.environ.setdefault("CF_CONTENT_ROOT", "/tmp/cf-smoke"))
# Keep the real ~/.config/content-forge/env well clear of the test.
os.environ.setdefault("CF_TOKEN_FILE", str(ROOT / "token-env"))
if ROOT.exists():
    shutil.rmtree(ROOT)

from fastapi.testclient import TestClient  # noqa: E402

from app import jobs, storage  # noqa: E402
from app.main import app  # noqa: E402

c = TestClient(app)
fails = []


def check(label, cond, detail=""):
    print(f"{'  ok ' if cond else 'FAIL '} {label}" + (f"  — {detail}" if detail and not cond else ""))
    if not cond:
        fails.append(label)


def settle(resp, timeout=25):
    """Follow a /jobs/<id>/wait redirect and block until the job settles."""
    job_id = resp.headers["location"].split("/")[2]
    for _ in range(timeout * 10):
        j = jobs.get(job_id)
        if j and j.settled:
            return j
        time.sleep(0.1)
    raise AssertionError("job never settled")


print("\n--- health ---")
h = c.get("/health")
check("GET /health", h.status_code == 200 and h.json()["ok"], h.text[:200])
check("health reports auth", "auth" in h.json() and h.json()["auth"]["method"],
      h.text[:200])

print("\n--- auth ---")
from app import auth as authmod  # noqa: E402

TOKFILE = Path(os.environ["CF_TOKEN_FILE"])
check("GET /auth", c.get("/auth").status_code == 200)
r = c.post("/auth/test")
check("POST /auth/test", r.status_code == 200 and "passed" in r.text, r.text[:200])

r = c.post("/auth/token", data={"token": "sk-ant-oat01-smoketestvalue0123456789"})
check("POST /auth/token", r.status_code == 200 and "Token saved" in r.text,
      r.text[:200])
check("token file written", TOKFILE.exists())
check("token file is 0600", oct(TOKFILE.stat().st_mode)[-3:] == "600",
      oct(TOKFILE.stat().st_mode))
check("token readable back", authmod.read_token()[0].endswith("0123456789"))
check("token reaches the CLI env",
      authmod.subprocess_env()["CLAUDE_CODE_OAUTH_TOKEN"].startswith("sk-ant-oat01"))
check("secret never rendered",
      "0123456789" not in c.get("/auth").text and "…" in authmod.mask("x" * 40))

age = authmod.token_age()
check("token save is stamped", age["exact"] is True, str(age))
check("token expiry projected ~1y out", 360 < age["days_left"] <= 365, str(age))
check("token not reported expired", age["expired"] is False)

# A stale fallback behind a working token is trivia, not an alarm: it belongs in
# the dormant section, and must not be what "in use" describes.
stale = ROOT / "stale-home" / ".credentials.json"
stale.parent.mkdir(parents=True, exist_ok=True)
stale.write_text(json.dumps({"claudeAiOauth": {"expiresAt": 1_000_000_000_000}}))
os.environ["CLAUDE_CONFIG_DIR"] = str(stale.parent)
check("token wins over a stale stored login",
      authmod.state()["method"] == "token"
      and authmod.state()["credentials_in_use"] is False)
page = c.get("/auth").text
check("stale login shown as not in use", "Not in use" in page and "dormant" in page)
check("healthy despite the stale fallback", c.get("/health").json()["ok"] is True)
os.environ.pop("CLAUDE_CONFIG_DIR")

os.environ["ANTHROPIC_API_KEY"] = "sk-ant-should-be-stripped"
check("api key stripped from CLI env",
      "ANTHROPIC_API_KEY" not in authmod.subprocess_env())
check("api key surfaced as a warning",
      "ANTHROPIC_API_KEY" in authmod.state()["stray_api_keys"])
os.environ.pop("ANTHROPIC_API_KEY")

check("empty token refused",
      "Paste a token" in c.post("/auth/token", data={"token": "  "}).text)
check("token survives a refused save", authmod.read_token()[0] != "")

r = c.post("/auth/token/delete")
check("POST /auth/token/delete", r.status_code == 200 and "removed" in r.text,
      r.text[:200])
check("token gone", authmod.read_token()[0] == "")

# an expired stored login must not read as healthy — the bug this page exists for
import app.claude_runner as cr  # noqa: E402
creds = ROOT / "fake-claude-home" / ".credentials.json"
creds.parent.mkdir(parents=True, exist_ok=True)
creds.write_text(json.dumps({"claudeAiOauth": {
    "expiresAt": 1_000_000_000_000, "subscriptionType": "max"}}))
os.environ["CLAUDE_CONFIG_DIR"] = str(creds.parent)
check("expired login detected", authmod.credentials_info()["expired"] is True)
cr.FAKE_CLAUDE = False
real_version = cr.cli_version
cr.cli_version = lambda: (True, "9.9.9 (stub)")
try:
    ok, note = cr.cli_available()
    check("expired login reported unhealthy", ok is False, f"{ok} {note}")
    check("expiry date named in the message", "2001-09" in note, note)
finally:
    cr.FAKE_CLAUDE = True
    cr.cli_version = real_version
    os.environ.pop("CLAUDE_CONFIG_DIR")

print("\n--- topics ---")
check("GET /", c.get("/").status_code == 200)
r = c.post("/topics", data={"name": "Youth Sports Parenting",
                            "description": "Club sports, burnout, money.",
                            "audience": "parents of 8-14 year olds"},
           follow_redirects=False)
check("POST /topics", r.status_code == 303, r.text[:200])
slug = r.headers["location"].split("/")[-1]
check("slug generated", slug == "youth-sports-parenting", slug)
check("topic folder exists", (ROOT / "topics" / slug / "articles").is_dir())
check("topic owns no voice folder", not (ROOT / "topics" / slug / "voice").exists())
for tab in ("articles", "write", "ideas", "settings"):
    check(f"GET topic?tab={tab}", c.get(f"/topics/{slug}?tab={tab}").status_code == 200)

print("\n--- voices ---")
check("GET /voices", c.get("/voices").status_code == 200)
r = c.post("/voices", data={"name": "Bob", "description": "Blunt, funny, short."},
           follow_redirects=False)
check("POST /voices", r.status_code == 303, r.text[:200])
vslug = r.headers["location"].split("/")[-1]
check("voice slug generated", vslug == "bob", vslug)
check("voice folder exists", (ROOT / "voices" / vslug / "samples").is_dir())
check("voice lives outside topics", not (ROOT / "topics" / slug / "voice").exists())

r = c.post(f"/voices/{vslug}/paste",
           data={"name": "tryouts night", "text": "Sample writing. " * 120},
           follow_redirects=False)
check("POST voices/paste", r.status_code == 303)
r = c.post(f"/voices/{vslug}/samples",
           files=[("files", ("second.md", b"More sample writing. " * 100, "text/markdown"))],
           follow_redirects=False)
check("POST voices/samples upload", r.status_code == 303)
check("2 samples on disk", len(storage.list_voice_samples(vslug)) == 2)

j = settle(c.post(f"/voices/{vslug}/build", data={"notes": "no exclamation marks"},
                  follow_redirects=False))
check("voice build job", j.status == "done", j.error)
check("profile.md written", storage.voice_profile_path(vslug).exists())
check("GET job status", c.get(f"/jobs/{j.id}/status").status_code == 200)

r = c.post(f"/voices/{vslug}/profile", data={"profile": "## Edited\nhand-tuned"},
           follow_redirects=False)
check("POST voices/profile save", r.status_code == 303
      and "hand-tuned" in storage.get_voice_profile(vslug))

# a second voice, to prove one topic can draw on many
r = c.post("/voices", data={"name": "Dana", "description": "Warm and lyrical."},
           follow_redirects=False)
vslug2 = r.headers["location"].split("/")[-1]
storage.save_voice_profile(vslug2, "## Dana\nLyrical and warm.")
check("second voice created", storage.get_voice(vslug2) is not None)
check("voices listed together", len(storage.list_voices()) == 2)
check("both voices offered on the write tab",
      vslug in c.get(f"/topics/{slug}?tab=write").text
      and vslug2 in c.get(f"/topics/{slug}?tab=write").text)

print("\n--- ideas ---")
j = settle(c.post(f"/topics/{slug}/ideas", data={"count": "2", "research": "1"},
                  follow_redirects=False))
check("ideas job", j.status == "done", j.error)
check("ideas persisted", len(storage.get_ideas(slug)) == 2)
check("ideas render", "Stub idea one" in c.get(f"/topics/{slug}?tab=ideas").text)

print("\n--- articles ---")
r = c.post(f"/topics/{slug}/articles",
           data={"title": "Why tryouts break kids", "prompt": "Make the case.",
                 "fmt": "opinion", "words": "800", "references": "1",
                 "voice": vslug, "styles": ["persuasive", "descriptive"]},
           follow_redirects=False)
check("POST articles", r.status_code == 303)
j = settle(r)
check("article job", j.status == "done", j.error)
aslug = j.result["article"]
a = storage.get_article(slug, aslug)
check("status ready", a.status == "ready", a.status)
check("voice recorded on the article", a.voice == vslug, a.voice)
check("styles recorded on the article",
      a.styles == ["persuasive", "descriptive"], str(a.styles))
check("unknown voice rejected",
      c.post(f"/topics/{slug}/articles",
             data={"title": "x", "prompt": "y", "voice": "ghost"},
             follow_redirects=False).status_code == 400)

# A voice with no profile cannot shape anything. Accepting it would generate
# generic prose while the article page claimed the voice was applied, so the
# web path must refuse it exactly as the MCP path does.
storage.create_voice("Unbuilt")
r = c.post(f"/topics/{slug}/articles",
           data={"title": "x", "prompt": "y", "voice": "unbuilt"},
           follow_redirects=False)
check("profile-less voice rejected on write", r.status_code == 400, r.text[:160])
check("rejection says why", "no profile" in r.text.lower(), r.text[:160])
check("profile-less voice rejected on rewrite",
      c.post(f"/topics/{slug}/articles/{aslug}/regenerate",
             data={"voice": "unbuilt"},
             follow_redirects=False).status_code == 400)
check("rejected rewrite left the article alone",
      storage.get_article(slug, aslug).voice == vslug)
check("profile-less voice not offered in the picker",
      "unbuilt" not in c.get(f"/topics/{slug}?tab=write").text)
check("title taken from H1", a.title == "Stub Article", a.title)
check("body on disk", "stub output" in storage.read_article_body(slug, aslug))
check("references included", "## References" in storage.read_article_body(slug, aslug))
check("GET article page", c.get(f"/topics/{slug}/articles/{aslug}").status_code == 200)
check("preview rendered server-side",
      "<h1>" in c.get(f"/topics/{slug}/articles/{aslug}").text)

r = c.post("/preview", data={"body": "# Hi\n\n- a\n- b"})
check("POST /preview", r.status_code == 200 and "<li>a</li>" in r.text, r.text[:120])

edited = "# Edited Title\n\nHand-edited **body** with a [link](https://example.com).\n"
r = c.post(f"/topics/{slug}/articles/{aslug}", data={"body": edited, "title": ""},
           follow_redirects=False)
check("POST save article", r.status_code == 303)
check("edit persisted", storage.get_article(slug, aslug).title == "Edited Title")

j = settle(c.post(f"/topics/{slug}/articles/{aslug}/regenerate",
                  data={"prompt": "different angle", "words": "1200",
                        "voice": vslug2, "styles": ["narrative"]},
                  follow_redirects=False))
check("regenerate job", j.status == "done", j.error)
check("word target updated", storage.get_article(slug, aslug).word_target == 1200)
check("references off when unchecked",
      storage.get_article(slug, aslug).with_references is False)
check("voice swapped on rewrite", storage.get_article(slug, aslug).voice == vslug2)
check("styles swapped on rewrite",
      storage.get_article(slug, aslug).styles == ["narrative"])

print("\n--- voice + style reach the prompt ---")
from app import prompts  # noqa: E402

sys_generic = prompts.article_system("", [])
check("no voice, no style stays generic",
      "VOICE PROFILE" not in sys_generic and "WRITING STYLE" not in sys_generic)

sys_voice = prompts.article_system("## Bob\nBlunt and short.", [])
check("voice profile embedded verbatim",
      "Blunt and short." in sys_voice and "AUTHOR VOICE PROFILE" in sys_voice)

sys_both = prompts.article_system("## Bob\nBlunt and short.",
                                  ["persuasive", "descriptive"])
check("both axes present at once",
      "AUTHOR VOICE PROFILE" in sys_both and "WRITING STYLE" in sys_both
      and "Persuasive" in sys_both and "Descriptive" in sys_both)
check("style survives without a voice",
      "WRITING STYLE" in prompts.article_system("", ["narrative"]))
check("unknown style ignored",
      "WRITING STYLE" not in prompts.article_system("", ["nonsense"]))

# the article carries its own voice, and it is that profile that gets loaded
picked = storage.get_article(slug, aslug)
check("article names the voice it was last written in", picked.voice == vslug2,
      picked.voice)
check("that voice's profile is what generation would load",
      storage.get_voice_profile(picked.voice).startswith("## Dana"),
      storage.get_voice_profile(picked.voice)[:40])

print("\n--- voice identity: rename and picture ---")
# A 1x1 PNG, and a 1x1 GIF, as real bytes.
PNG = bytes.fromhex(
    "89504e470d0a1a0a0000000d49484452000000010000000108060000001f15c489"
    "0000000a49444154789c63000100000500010d0a2db40000000049454e44ae426082")
GIF = b"GIF89a\x01\x00\x01\x00\x80\x00\x00\x00\x00\x00\xff\xff\xff!" \
      b"\xf9\x04\x01\x00\x00\x00\x00,\x00\x00\x00\x00\x01\x00\x01\x00" \
      b"\x00\x02\x02D\x01\x00;"

check("no picture initially", storage.get_voice(vslug).has_image is False)
r = c.post(f"/voices/{vslug}/image",
           files=[("image", ("me.png", PNG, "image/png"))],
           follow_redirects=False)
check("POST voice image", r.status_code == 303, r.text[:200])
check("picture recorded", storage.get_voice(vslug).has_image is True)
r = c.get(f"/voices/{vslug}/image")
check("GET voice image", r.status_code == 200 and r.content == PNG)
check("served as png", r.headers["content-type"] == "image/png",
      r.headers.get("content-type", ""))
check("served with nosniff",
      r.headers.get("x-content-type-options") == "nosniff")

# Replacing swaps the file rather than accumulating one per format.
c.post(f"/voices/{vslug}/image", files=[("image", ("me.gif", GIF, "image/gif"))],
       follow_redirects=False)
check("replacing leaves exactly one portrait",
      len([f for f in storage.voice_dir(vslug).iterdir()
           if f.name.startswith("portrait")]) == 1)
check("replacement served as gif",
      c.get(f"/voices/{vslug}/image").headers["content-type"] == "image/gif")

# Content type is decided by the bytes, not the filename an uploader chose.
r = c.post(f"/voices/{vslug}/image",
           files=[("image", ("evil.png", b"<script>alert(1)</script>", "image/png"))],
           follow_redirects=False)
check("non-image rejected despite .png name", r.status_code == 400, r.text[:160])
check("rejection did not disturb the existing picture",
      c.get(f"/voices/{vslug}/image").content == GIF)
r = c.post(f"/voices/{vslug}/image",
           files=[("image", ("big.png", PNG + b"\0" * (6 * 1024 * 1024), "image/png"))],
           follow_redirects=False)
check("oversized image rejected", r.status_code == 400, r.text[:160])

check("picture shows on the voices index",
      f"/voices/{vslug}/image" in c.get("/voices").text)
check("voice without a picture falls back to initials",
      "avatar-fallback" in c.get(f"/voices/{vslug2}").text)

r = c.post(f"/voices/{vslug}/image/delete", follow_redirects=False)
check("DELETE voice image", r.status_code == 303
      and storage.get_voice(vslug).has_image is False)
check("image 404s once removed", c.get(f"/voices/{vslug}/image").status_code == 404)

# Renaming must not break the link from articles, which is by slug.
before = storage.get_article(slug, aslug).voice
r = c.post(f"/voices/{vslug2}/edit",
           data={"name": "Renamed Dana", "description": "Still lyrical."},
           follow_redirects=False)
check("POST voice rename", r.status_code == 303)
check("name changed", storage.get_voice(vslug2).name == "Renamed Dana")
check("slug unchanged by rename", storage.get_voice(vslug2).slug == vslug2)
check("article still linked after rename",
      storage.get_article(slug, aslug).voice == before)
check("new name renders on the article",
      "Renamed Dana" in c.get(f"/topics/{slug}/articles/{aslug}").text)

print("\n--- exports ---")
r = c.get(f"/topics/{slug}/articles/{aslug}/export.docx")
check("GET export.docx", r.status_code == 200 and len(r.content) > 5000,
      f"{r.status_code} {len(r.content)}b")
check("docx content-type", "wordprocessingml" in r.headers["content-type"])
r = c.get(f"/topics/{slug}/articles/{aslug}/raw.md")
check("GET raw.md", r.status_code == 200)

# open the docx for real
import docx  # noqa: E402
d = docx.Document(str(storage.article_dir(slug, aslug) / "export.docx"))
paras = [p.text for p in d.paragraphs]
styles = {p.style.name for p in d.paragraphs}
check("docx parses", len(paras) > 3, str(len(paras)))
check("docx has a title", any(s in styles for s in ("Title", "Heading 1")), str(styles))
check("docx has heading", "Heading 2" in styles or "Heading 1" in styles, str(styles))
check("docx has bullets", "List Bullet" in styles, str(styles))
check("docx has numbers", "List Number" in styles, str(styles))
check("docx has quote", any("Quote" in s for s in styles), str(styles))
bold = [r.text for p in d.paragraphs for r in p.runs if r.bold]
ital = [r.text for p in d.paragraphs for r in p.runs if r.italic]
check("inline bold survived", "bold" in bold, str(bold))
check("inline italic survived", "italic" in ital, str(ital))
xml = d.element.xml
check("hyperlink survived", "hyperlink" in xml)
check("markdown syntax stripped",
      "**" not in "\n".join(paras) and "](" not in "\n".join(paras))

print("\n--- deletes ---")
check("delete article", c.post(f"/topics/{slug}/articles/{aslug}/delete",
                               follow_redirects=False).status_code == 303
      and not storage.article_dir(slug, aslug).exists())
check("delete sample", c.post(f"/voices/{vslug}/samples/delete",
                              data={"name": "second.md"},
                              follow_redirects=False).status_code == 303
      and len(storage.list_voice_samples(vslug)) == 1)
check("delete topic", c.post(f"/topics/{slug}/delete",
                             follow_redirects=False).status_code == 303
      and not storage.topic_dir(slug).exists())
check("voices survive their topic",
      storage.get_voice(vslug) is not None and storage.get_voice(vslug2) is not None)
check("delete voice", c.post(f"/voices/{vslug}/delete",
                             follow_redirects=False).status_code == 303
      and not storage.voice_dir(vslug).exists())

print("\n--- errors ---")
check("404 unknown topic", c.get("/topics/nope").status_code == 404)
check("404 unknown voice", c.get("/voices/nope").status_code == 404)
check("404 unknown job", c.get("/jobs/deadbeef/wait").status_code == 404)

print()
if fails:
    print(f"{len(fails)} FAILURE(S): " + ", ".join(fails))
    sys.exit(1)
print("ALL CHECKS PASSED")
